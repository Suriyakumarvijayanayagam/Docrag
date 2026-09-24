"""
Tables in technical documents (spec sheets, pinouts, parameter tables) are
where exact answers live. Chunk-and-embed treats a table row as fuzzy text,
which is wrong for "what's the max operating temperature" - that needs a
lookup, not a semantic guess.

Tables are extracted at ingest into structured_facts as label/value pairs and
matched against the question before retrieval. Matches are shown to the user
as exact sources, so a false positive is expensive: matching is whole-word and
only the strongest overlap tier is kept.
"""
from dataclasses import dataclass
from pathlib import Path

from app.text_match import content_tokens


@dataclass
class Fact:
    page: int | None
    locator: str
    label: str
    value: str


def _rows_to_facts(rows: list[list[str]], locator: str, page: int | None) -> list[Fact]:
    """
    Two-column tables are label/value pairs. Wider tables use the header row
    as column labels, flattened to "RowEntity.ColumnHeader -> cell".
    """
    rows = [[" ".join(cell.split()) for cell in row] for row in rows]
    rows = [row for row in rows if any(row)]
    facts: list[Fact] = []
    if not rows:
        return facts
    if all(len(row) == 2 for row in rows):
        for index, (label, value) in enumerate(rows, start=1):
            if label and value:
                facts.append(Fact(page, f"{locator} row {index}", label, value))
    elif len(rows) > 1:
        header = rows[0]
        for index, row in enumerate(rows[1:], start=2):
            entity = row[0] or f"row {index}"
            for column, value in enumerate(row[1:], start=1):
                if column < len(header) and header[column] and value:
                    facts.append(Fact(page, f"{locator} row {index}", f"{entity}.{header[column]}", value))
    return facts


def extract_facts(path: str) -> list[Fact]:
    ext = Path(path).suffix.lower()
    facts: list[Fact] = []
    if ext == ".pdf":
        import fitz
        with fitz.open(path) as pdf:
            for page_number, page in enumerate(pdf, start=1):
                try:
                    tables = page.find_tables().tables
                except Exception:
                    continue
                for table_number, table in enumerate(tables, start=1):
                    rows = [[str(cell).strip() if cell else "" for cell in row] for row in table.extract()]
                    facts.extend(_rows_to_facts(rows, f"p. {page_number} table {table_number}", page_number))
    elif ext == ".docx":
        from docx import Document
        for table_number, table in enumerate(Document(path).tables, start=1):
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            facts.extend(_rows_to_facts(rows, f"table {table_number}", None))
    elif ext == ".pptx":
        from pptx import Presentation
        for slide_number, slide in enumerate(Presentation(path).slides, start=1):
            for shape in slide.shapes:
                if getattr(shape, "has_table", False):
                    rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                    facts.extend(_rows_to_facts(rows, f"slide {slide_number} table", slide_number))
    return facts


def match_facts(question: str, facts: list[dict], limit: int = 8) -> list[dict]:
    """
    Keeps facts whose label shares the most whole-word tokens with the
    question. A question matching one label on two words and another on one is
    asking about the first, so only the top overlap tier survives.
    """
    tokens = content_tokens(question)
    if not tokens:
        return []
    scored = []
    for fact in facts:
        overlap = len(tokens & content_tokens(fact["label"]))
        if overlap:
            scored.append((overlap, fact))
    if not scored:
        return []
    best = max(score for score, _ in scored)
    return [fact for score, fact in scored if score == best][:limit]
