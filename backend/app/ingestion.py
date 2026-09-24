import io
import re
from pathlib import Path

import fitz
import pytesseract
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation

from app.chunking import TextBlock, chunk_blocks
from app.config import settings

SUPPORTED_EXTENSIONS = {
    ".pdf", ".docx", ".txt", ".md", ".pptx", ".xlsx", ".html", ".htm",
    ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp",
}


def _looks_numeric(value: str) -> bool:
    return bool(re.fullmatch(r"[\s$€£₹%.,:/-]*\d[\d\s$€£₹%.,:/-]*", value))


def extract_document(path: str, extension: str) -> tuple[list[TextBlock], int]:
    ext = extension.lower()
    blocks: list[TextBlock] = []
    page_count = 0
    if ext in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}:
        text = pytesseract.image_to_string(Image.open(path)).strip()
        if text:
            blocks.append(TextBlock(text, 1))
        page_count = 1
    elif ext == ".pdf":
        pdf = fitz.open(path)
        page_count = len(pdf)
        for page_index, page in enumerate(pdf, start=1):
            text = page.get_text("text").strip()
            if len(text) < 24:
                pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                image = Image.open(io.BytesIO(pixmap.tobytes("png")))
                text = pytesseract.image_to_string(image).strip()
            if text:
                blocks.append(TextBlock(text, page_index))
    elif ext == ".docx":
        document = DocxDocument(path)
        page_count = 0
        section = ""
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style = paragraph.style.name if paragraph.style else ""
            if style.lower().startswith("heading"):
                section = text
            blocks.append(TextBlock(text, None, section))
        for table_index, table in enumerate(document.tables, start=1):
            for row in table.rows:
                values = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(values):
                    blocks.append(TextBlock(" | ".join(values), None, f"Table {table_index}"))
    elif ext == ".pptx":
        presentation = Presentation(path)
        page_count = len(presentation.slides)
        for slide_number, slide in enumerate(presentation.slides, start=1):
            title = ""
            if slide.shapes.title and slide.shapes.title.text:
                title = slide.shapes.title.text.strip()
            lines = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and shape.text.strip():
                    lines.append(shape.text.strip())
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        lines.append(" | ".join(cell.text.strip() for cell in row.cells))
            blocks.append(TextBlock("\n".join(dict.fromkeys(lines)), slide_number, title))
    elif ext == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            rows = [
                (row_number, list(row))
                for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1)
            ]
            if not rows:
                continue

            # Workbooks often have a title/date line before the actual header row.
            # Pick the early row with the most text labels (headers are words,
            # data rows usually mix in numbers), then the densest, then the earliest.
            header_candidates = [
                (row_number, row)
                for row_number, row in rows[:20]
                if any(value is not None and str(value).strip() for value in row)
            ]
            if not header_candidates:
                continue

            def header_score(candidate: tuple[int, list]) -> tuple[int, int, int]:
                row_number, row = candidate
                values = [value for value in row if value is not None and str(value).strip()]
                text_labels = sum(isinstance(value, str) and not _looks_numeric(value) for value in values)
                return (text_labels, len(values), -row_number)

            header_row_number, header_row = max(header_candidates, key=header_score)
            headers = [
                str(value).strip() if value is not None and str(value).strip() else f"Column {i + 1}"
                for i, value in enumerate(header_row)
            ]
            for row_number, row in rows:
                if row_number <= header_row_number:
                    continue
                cells = [
                    f"{headers[i]}: {value}"
                    for i, value in enumerate(row)
                    if value is not None and str(value).strip() and i < len(headers)
                ]
                if cells:
                    blocks.append(TextBlock(f"Sheet {sheet.title} — Row {row_number} — " + " | ".join(cells), None, sheet.title))
        page_count = len(workbook.worksheets)
    elif ext in {".html", ".htm"}:
        soup = BeautifulSoup(Path(path).read_text(errors="ignore"), "html.parser")
        for node in soup(["script", "style", "nav", "footer"]):
            node.decompose()
        for element in soup.find_all(["h1", "h2", "h3", "p", "li", "td", "th"]):
            text = element.get_text(" ", strip=True)
            if text:
                blocks.append(TextBlock(text, None, "Heading" if element.name.startswith("h") else ""))
    else:
        text = Path(path).read_text(errors="replace")
        section = ""
        for line in text.splitlines():
            line = line.strip()
            if not line:
                continue
            if ext == ".md" and line.startswith("#"):
                section = line.lstrip("# ").strip()
            blocks.append(TextBlock(line, None, section))

    if not any(block.text.strip() for block in blocks):
        raise ValueError("No readable text was extracted. The document may be empty, encrypted, or image-only.")
    return blocks, page_count


def make_chunks(path: str) -> tuple[list, int]:
    blocks, pages = extract_document(path, Path(path).suffix)
    # Spreadsheet rows are self-contained records (each carries its column labels).
    # Overlap duplicates records across chunks and can make count/list questions ambiguous.
    overlap = 0 if Path(path).suffix.lower() == ".xlsx" else settings.chunk_overlap
    chunks = chunk_blocks(blocks, settings.chunk_tokens, overlap)
    if not chunks:
        raise ValueError("Document did not produce any searchable text chunks")
    return chunks, pages
