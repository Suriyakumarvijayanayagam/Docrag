"""
Regenerates the sample documents in ./sample-docs.

They are deliberately built to exercise every path in the pipeline: ruled
tables (structured-fact extraction), an embedded block diagram (figure
extraction), multi-page prose (deep retrieval), two datasheets that disagree
(cross-document comparison), and a requirements doc whose limits conflict with
one of them. Nothing in them mentions CAN bus, Modbus, or USB - those make
good "does it admit what it doesn't know" probes.

    cd backend && python -m scripts.make_sample_docs
"""
import os

import fitz
from docx import Document
from docx.shared import Pt
from pathlib import Path

OUT = Path(os.environ.get("SAMPLE_DOCS_DIR") or Path(__file__).resolve().parents[2] / "sample-docs")
OUT.mkdir(parents=True, exist_ok=True)


def block_diagram_png(path):
    """Render a real block diagram to PNG so figure extraction has something true to find."""
    d = fitz.open()
    p = d.new_page(width=420, height=260)
    boxes = [(20, 60, 110, 110, "24V INPUT"), (150, 60, 250, 110, "PC-42 LDO"),
             (290, 60, 400, 110, "3V3 LOGIC RAIL"), (150, 160, 250, 210, "THERMAL\nSHUTDOWN")]
    for x0, y0, x1, y1, label in boxes:
        p.draw_rect(fitz.Rect(x0, y0, x1, y1), color=(0.1, 0.2, 0.5), width=1.5)
        for i, line in enumerate(label.split("\n")):
            p.insert_text((x0 + 8, y0 + 22 + i * 12), line, fontsize=8)
    for a, b in [((110, 85), (150, 85)), ((250, 85), (290, 85)), ((200, 110), (200, 160))]:
        p.draw_line(fitz.Point(*a), fitz.Point(*b), color=(0, 0, 0), width=1.2)
    p.insert_text((20, 240), "Figure 1: PC-42 power subsystem block diagram", fontsize=9)
    pix = p.get_pixmap(matrix=fitz.Matrix(2, 2))
    pix.save(path)
    d.close()


def ruled_table(page, x0, y0, rows, col_w=(230, 150)):
    """Draw a ruled table so PyMuPDF's find_tables() detects it."""
    rh = 22
    for r, row in enumerate(rows):
        y = y0 + r * rh
        x = x0
        for c, cell in enumerate(row):
            w = col_w[c]
            page.draw_rect(fitz.Rect(x, y, x + w, y + rh), color=(0.3, 0.3, 0.3), width=0.8)
            page.insert_text((x + 6, y + 15), cell, fontsize=9,
                             fontname="hebo" if r == 0 else "helv")
            x += w
    return y0 + len(rows) * rh


# ---------- Datasheet A: PC-42 LDO regulator (PDF, 3 pages, table + figure) ----------
doc = fitz.open()
p = doc.new_page()
p.insert_text((60, 70), "PowerCell PC-42", fontsize=22)
p.insert_text((60, 95), "Low-Dropout Linear Regulator - Datasheet Rev C", fontsize=11)
p.insert_textbox(fitz.Rect(60, 120, 540, 260),
    "The PowerCell PC-42 is a low-dropout linear regulator designed for noise-sensitive "
    "analog and mixed-signal rails in industrial control equipment. It accepts a wide "
    "6 V to 24 V input and delivers a fixed 3.3 V output at up to 1.5 A continuous. "
    "Dropout voltage is 310 mV at full load. The device integrates thermal shutdown, "
    "current limiting, and reverse-polarity protection.\n\n"
    "Because the PC-42 is a linear topology, power dissipation is the input-to-output "
    "differential multiplied by load current. At 24 V input and 1.5 A load the device "
    "dissipates 31 W, which exceeds the package rating without a heatsink; the "
    "recommended maximum input for continuous full-load operation is 12 V.", fontsize=10)
p.insert_text((60, 300), "Table 1: Absolute Maximum Ratings", fontsize=11)
ruled_table(p, 60, 315, [
    ("Parameter", "Value"),
    ("Input voltage range", "6 V to 24 V"),
    ("Output voltage", "3.3 V fixed"),
    ("Continuous output current", "1.5 A"),
    ("Dropout voltage at 1.5 A", "310 mV"),
    ("Thermal shutdown threshold", "150 C"),
    ("Quiescent current", "45 uA"),
    ("Operating temperature range", "-40 C to +125 C"),
])

p2 = doc.new_page()
p2.insert_text((60, 70), "Figure 1: PC-42 power subsystem block diagram", fontsize=11)
png = OUT / ".blockdiagram.tmp.png"
block_diagram_png(str(png))
p2.insert_image(fitz.Rect(60, 90, 480, 350), filename=str(png))
p2.insert_textbox(fitz.Rect(60, 370, 540, 520),
    "The 24 V bulk rail enters through a reverse-polarity FET and is regulated down to "
    "the 3V3 logic rail. The thermal shutdown block monitors junction temperature and "
    "disables the pass element at 150 C, re-enabling with 15 C of hysteresis.\n\n"
    "Efficiency at nominal 12 V input and 3.3 V output is 27 percent, which is inherent "
    "to linear regulation and not a defect of this part.", fontsize=10)

p3 = doc.new_page()
p3.insert_text((60, 70), "Application Notes", fontsize=14)
p3.insert_textbox(fitz.Rect(60, 95, 540, 400),
    "Output capacitor selection: the PC-42 requires a minimum 10 uF ceramic output "
    "capacitor with ESR below 100 mOhm for stability. X7R dielectric is recommended "
    "over Y5V because of Y5V's bias-voltage derating.\n\n"
    "Layout: keep the feedback trace away from the switching node of any upstream "
    "converter. The exposed pad must be tied to the ground plane with at least nine "
    "thermal vias.\n\n"
    "Qualification: the PC-42 is AEC-Q100 Grade 1 qualified and carries a 10-year "
    "supply commitment from the date of Rev C release.", fontsize=10)
doc.save(str(OUT / "PC-42_datasheet.pdf"))
doc.close()

# ---------- Datasheet B: VE-9 switching regulator (PDF, 2 pages, table) ----------
doc = fitz.open()
p = doc.new_page()
p.insert_text((60, 70), "VoltEdge VE-9", fontsize=22)
p.insert_text((60, 95), "Synchronous Buck Regulator - Datasheet Rev A", fontsize=11)
p.insert_textbox(fitz.Rect(60, 120, 540, 250),
    "The VoltEdge VE-9 is a synchronous buck converter targeting the same 3.3 V logic "
    "rails as traditional linear parts, but at substantially higher efficiency. It "
    "accepts 4.5 V to 36 V input and delivers up to 3 A continuous at 94 percent peak "
    "efficiency, switching at 500 kHz.\n\n"
    "Unlike a linear regulator, the VE-9 injects switching ripple onto the output rail: "
    "35 mV peak-to-peak typical at full load. Noise-sensitive analog sections may "
    "require post-regulation or additional filtering.", fontsize=10)
p.insert_text((60, 290), "Table 1: Electrical Characteristics", fontsize=11)
ruled_table(p, 60, 305, [
    ("Parameter", "Value"),
    ("Input voltage range", "4.5 V to 36 V"),
    ("Output voltage", "3.3 V adjustable"),
    ("Continuous output current", "3 A"),
    ("Peak efficiency", "94 percent"),
    ("Switching frequency", "500 kHz"),
    ("Output ripple at full load", "35 mV pp"),
    ("Thermal shutdown threshold", "165 C"),
    ("Quiescent current", "12 uA"),
])
p2 = doc.new_page()
p2.insert_text((60, 70), "Reliability and Supply", fontsize=14)
p2.insert_textbox(fitz.Rect(60, 95, 540, 320),
    "The VE-9 is qualified to AEC-Q100 Grade 2 only, and is not rated for the -40 C to "
    "+125 C ambient range required by Grade 1 designs. Sustained operation above +105 C "
    "ambient is outside the datasheet envelope.\n\n"
    "The VE-9 entered volume production in Q1 and carries a 5-year supply commitment. "
    "Lead time is currently 22 weeks.", fontsize=10)
doc.save(str(OUT / "VE-9_datasheet.pdf"))
doc.close()

# ---------- DOCX: system requirements with tables ----------
d = Document()
d.add_heading("Orion Controller - Power Subsystem Requirements", level=1)
d.add_paragraph("This document captures the binding requirements for the Orion "
                "industrial controller power subsystem. Any regulator selected for "
                "the 3.3 V logic rail must satisfy every requirement in Table 1.")
d.add_heading("Table 1: Binding Requirements", level=2)
t = d.add_table(rows=0, cols=2)
for label, value in [
    ("Requirement ID prefix", "ORN-PWR"),
    ("Logic rail voltage", "3.3 V plus or minus 2 percent"),
    ("Minimum continuous current", "2.0 A"),
    ("Maximum output ripple", "20 mV pp"),
    ("Required qualification", "AEC-Q100 Grade 1"),
    ("Required supply commitment", "10 years minimum"),
    ("Maximum lead time", "16 weeks"),
    ("Ambient temperature range", "-40 C to +125 C"),
]:
    row = t.add_row()
    row.cells[0].text, row.cells[1].text = label, value
d.add_heading("Notes", level=2)
d.add_paragraph("The 20 mV ripple limit is driven by the 16-bit ADC on the sensor "
                "board. Requirement ORN-PWR-014 states that any switching regulator "
                "must be followed by post-regulation if this limit cannot be met "
                "directly.")
d.save(str(OUT / "Orion_requirements.docx"))

# A welding procedure specification. Fictional: made up for testing the
# pipeline, not a qualified procedure. The wide parameter table exercises the
# "entity.column" flattening in fact extraction.
d = Document()
d.add_heading("WPS-SMAW-017 - Welding Procedure Specification", level=1)
d.add_paragraph("SAMPLE DOCUMENT FOR TESTING ONLY. Fictional company and values; not a qualified procedure.")
d.add_paragraph("Scope: single-V butt welds in structural steel plate for general fabrication, "
                "supported by PQR-017. Welders must hold a current qualification for the positions listed.")
d.add_heading("Table 1: Procedure variables", level=2)
t = d.add_table(rows=0, cols=2)
for label, value in [
    ("Welding process", "SMAW (manual metal arc)"),
    ("Joint design", "Single-V butt, 60 degree included angle, 2 mm root gap, 1.5 mm root face"),
    ("Base metal", "IS 2062 E250 BR"),
    ("Thickness range qualified", "10 mm to 25 mm"),
    ("Filler metal", "AWS A5.1 E7018, 3.15 mm and 4.0 mm"),
    ("Positions", "1G (PA) and 2G (PC)"),
    ("Minimum preheat", "50 C"),
    ("Maximum interpass temperature", "250 C"),
    ("Post-weld heat treatment", "None"),
    ("Electrode baking", "300 to 350 C for 2 hours, then held at 100 to 150 C in a holding oven"),
]:
    row = t.add_row()
    row.cells[0].text, row.cells[1].text = label, value
d.add_heading("Table 2: Welding parameters", level=2)
t = d.add_table(rows=1, cols=5)
for cell, text in zip(t.rows[0].cells, ["Pass", "Electrode (mm)", "Current (A)", "Voltage (V)", "Travel speed (mm/min)"]):
    cell.text = text
for values in [("Root", "3.15", "90-110", "21-23", "80-100"),
               ("Fill", "4.0", "140-170", "22-25", "120-160"),
               ("Cap", "4.0", "140-165", "22-24", "130-170")]:
    row = t.add_row()
    for cell, text in zip(row.cells, values):
        cell.text = text
d.add_heading("Inspection", level=2)
d.add_paragraph("All welds receive 100 percent visual inspection. The root pass is back-gouged to sound "
                "metal before the second-side weld. Arc strikes outside the weld groove are not permitted.")
d.save(str(OUT / "Sample_WPS-SMAW-017.docx"))

# ---------- Sample job file JOB-2025-118 (fictional, with planted mistakes) ----------
# Everything needed for the job checks, with problems a QA engineer would want
# caught before an inspector finds them:
#   PQR coupon 10 mm, WPS claims up to 25 mm (over 2T)
#   consumable certificate is for E7016, WPS specifies E7018
#   weld log: W-14 welded 3G (qualified 1G/2G), W-22 is GMAW-only on an SMAW
#   WPS, W-31 has no qualification, a joint cites WPS-SMAW-021 (not in the job),
#   a 32 mm joint on a 10-25 mm WPS, and W-14 idle > 6 months before J-104.
JOB = OUT / "job-2025-118"
JOB.mkdir(exist_ok=True)
d.save(str(JOB / "Sample_WPS-SMAW-017.docx"))


def record_pdf(path, title, subtitle, rows, note):
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((60, 70), title, fontsize=16)
    page.insert_text((60, 92), subtitle, fontsize=10)
    page.insert_text((60, 110), "SAMPLE DOCUMENT FOR TESTING ONLY - fictional company, people and values.", fontsize=8)
    bottom = ruled_table(page, 60, 130, rows, col_w=(200, 260))
    page.insert_text((60, bottom + 30), note, fontsize=9)
    doc.save(str(path))


record_pdf(JOB / "Sample_PQR-017.pdf", "PQR-017", "Procedure Qualification Record", [
    ("Item", "Record"),
    ("PQR No.", "PQR-017"),
    ("Welding process", "SMAW"),
    ("Base metal", "IS 2062 E250 BR"),
    ("Coupon thickness", "10 mm"),
    ("Filler metal", "AWS A5.1 E7018, 3.15 mm and 4.0 mm"),
    ("Test position", "2G (PC)"),
    ("Preheat", "50 C"),
    ("Date of test", "12/03/2025"),
    ("Tensile and bend tests", "Acceptable"),
], "Supports WPS-SMAW-017.")

for welder_id, name, process, tested_in, qualified, filler, tested_on in (
    ("W-14", "R. Kumar", "SMAW", "2G (PC)", "1G, 2G", "AWS A5.1 E7018", "20/01/2025"),
    ("W-22", "A. Singh", "GMAW", "3G (PF)", "1G, 2G, 3G, 4G", "AWS A5.18 ER70S-6", "01/02/2025"),
):
    record_pdf(JOB / f"Sample_WPQ-{welder_id.replace('-', '')}.pdf", f"WPQ {welder_id}", "Welder Performance Qualification Record", [
        ("Item", "Record"),
        ("Welder name", name),
        ("Welder ID", welder_id),
        ("Welding process", process),
        ("Test position", tested_in),
        ("Positions qualified", qualified),
        ("Filler metal", filler),
        ("Date of test", tested_on),
        ("Result", "Qualified"),
    ], "Visual and bend tests witnessed by the QA engineer.")

record_pdf(JOB / "Sample_TC-4471.pdf", "Test Certificate TC-4471", "Inspection certificate EN 10204 3.1 - covered electrodes", [
    ("Item", "Record"),
    ("Product", "Low-hydrogen covered electrode"),
    ("AWS classification", "A5.1 E7016"),
    ("Batch no.", "4471"),
    ("Diameter", "4.0 mm"),
    ("Date of issue", "05/03/2025"),
], "Chemical and mechanical results conform to the classification.")

from datetime import datetime as _dt
from openpyxl import Workbook

book = Workbook()
sheet = book.active
sheet.title = "Weld log"
sheet.append(["JOB-2025-118 pipe rack - weld log (SAMPLE, fictional)"])
sheet.append([])
sheet.append(["Joint No.", "Welder ID", "WPS No.", "Date welded", "Position", "Thickness (mm)"])
for joint, welder, wps, day, position, thickness in (
    ("J-101", "W-14", "WPS-SMAW-017", "2025-04-02", "1G", 12),
    ("J-102", "W-14", "WPS-SMAW-017", "2025-04-20", "3G", 12),
    ("J-103", "W-22", "WPS-SMAW-017", "2025-04-21", "1G", 12),
    ("J-105", "W-31", "WPS-SMAW-017", "2025-05-01", "1G", 12),
    ("J-106", "W-14", "WPS-SMAW-021", "2025-05-02", "1G", 12),
    ("J-107", "W-14", "WPS-SMAW-017", "2025-05-03", "1G", 32),
    ("J-104", "W-14", "WPS-SMAW-017", "2025-11-15", "2G", 16),
):
    sheet.append([joint, welder, wps, _dt.fromisoformat(day), position, thickness])
book.save(str(JOB / "Sample_WeldLog_JOB-2025-118.xlsx"))

png.unlink(missing_ok=True)  # intermediate used to build the PDF, not a sample

for f in sorted(p for p in OUT.rglob("*") if p.is_file()):
    print(f"  {f.name}  ({f.stat().st_size//1024} KB)")
