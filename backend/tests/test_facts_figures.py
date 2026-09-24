from app.facts import _rows_to_facts, extract_facts, match_facts
from app.figures import best_figure
from app.text_match import content_tokens


def test_two_column_tables_are_label_value_pairs():
    facts = _rows_to_facts([["Input voltage", "4.5 - 40 V"], ["Dropout", "300 mV"], ["", ""]], "table 1", None)

    assert [(f.label, f.value) for f in facts] == [("Input voltage", "4.5 - 40 V"), ("Dropout", "300 mV")]


def test_wide_tables_flatten_to_entity_dot_column():
    rows = [["Parameter", "Min", "Max"], ["Vin", "4.5 V", "40 V"], ["Tj", "", "150 C"]]

    facts = _rows_to_facts(rows, "p. 2 table 1", 2)

    assert [(f.label, f.value) for f in facts] == [("Vin.Min", "4.5 V"), ("Vin.Max", "40 V"), ("Tj.Max", "150 C")]
    assert all(f.page == 2 for f in facts)


def test_docx_tables_are_extracted(tmp_path):
    from docx import Document

    path = tmp_path / "spec.docx"
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Operating temperature", "-40 to 125 C"
    table.cell(1, 0).text, table.cell(1, 1).text = "Package", "SOT-23"
    document.save(path)

    facts = extract_facts(str(path))

    assert ("Operating temperature", "-40 to 125 C") in [(f.label, f.value) for f in facts]


def test_fact_matching_is_whole_word_and_keeps_strongest_tier():
    facts = [
        {"label": "Thermal shutdown temperature"},
        {"label": "Operating temperature"},
        {"label": "Part number"},
    ]

    matched = match_facts("What is the thermal shutdown temperature?", facts)

    assert [f["label"] for f in matched] == ["Thermal shutdown temperature"]
    assert match_facts("What is the art direction?", facts) == []


def test_stopwords_and_substrings_do_not_match():
    assert "the" not in content_tokens("the thermal part")
    assert content_tokens("PC-42 pinout") == {"pc-42", "pinout"}


def test_figure_matching_is_conservative():
    figures = [
        {"id": 1, "caption": "Figure 1. PC-42 block diagram: 24V input to 3V3 logic rail"},
        {"id": 2, "caption": "Image from page 5"},
    ]

    assert best_figure("diagram: PC-42 block diagram of the logic rail", figures)["id"] == 1
    assert best_figure("diagram of the CAN bus transceiver", figures) is None
    assert best_figure("diagram", figures) is None
