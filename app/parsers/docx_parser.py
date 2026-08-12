"""
DOCX text extraction. DOCX has no native "page" concept (pagination is a
rendering-time thing), so we chunk by section/paragraph-block index instead
and label it clearly for the user.
"""
from docx import Document


def parse_docx(file_path: str) -> list[dict]:
    """
    Returns: [{"section": 1, "text": "..."}, ...]
    Groups paragraphs into blocks of ~20 to keep parity with PDF page granularity.
    """
    doc = Document(file_path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    # also pull table content, tables are often where the real info lives
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))

    blocks = []
    block_size = 20
    for i in range(0, len(paragraphs), block_size):
        chunk_text = "\n".join(paragraphs[i:i + block_size])
        if chunk_text.strip():
            blocks.append({"section": (i // block_size) + 1, "text": chunk_text})
    return blocks
